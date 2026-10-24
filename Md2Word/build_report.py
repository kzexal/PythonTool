"""Build a thesis DOCX from Markdown while reusing a Word template's styles.

Supported Markdown: headings, paragraphs, **bold**, *italic*, `inline code`, links,
ordered/unordered lists, fenced code, images, pipe tables, citations [@key].
Image caption: ![Caption text](relative/path.png) — Word Caption + SEQ Hình
Table caption: a line `Table: Caption text` immediately before a Markdown table — Word Caption + SEQ Bảng.
"""
from __future__ import annotations
import argparse, json, re, sys
from pathlib import Path

if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HEADING_RE = re.compile(r'^(#{1,9})\s+(.+?)\s*#*\s*$')
IMAGE_RE = re.compile(r'^!\[(.*?)\]\((.+?)\)\s*$')
CITATION_RE = re.compile(r'\[@([A-Za-z0-9_.:-]+)\]')
LINK_RE = re.compile(r'\[([^\]]+)\]\((https?://[^)]+)\)')


def disable_numbering(paragraph):
    """Disable list numbering on a paragraph even if its style has numPr."""
    pPr = paragraph._p.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is None:
        numPr = OxmlElement('w:numPr')
        pPr.append(numPr)
    else:
        numPr.clear()
    ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), '0')
    numId = OxmlElement('w:numId'); numId.set(qn('w:val'), '0')
    numPr.append(ilvl); numPr.append(numId)


def template_has_heading_numbering(doc):
    """Check if Heading 1 in the template has automatic multilevel numbering."""
    try:
        s = doc.styles['Heading 1']
        pPr = s._element.pPr
        if pPr is not None and pPr.numPr is not None:
            numId = pPr.numPr.find(qn('w:numId'))
            if numId is not None and numId.get(qn('w:val')) != '0':
                return True
    except Exception:
        pass
    return False


def clean_heading_text(title, level):
    """Strip chapter/section numbers so Word's automatic numbering does not duplicate them."""
    if level == 1:
        cleaned = re.sub(r'^\s*CHƯƠNG\s+\d+[\.:\s-]*', '', title, flags=re.IGNORECASE)
        if cleaned == title:
            cleaned = re.sub(r'^\s*\d+[\.:\s-]+', '', title)
        return cleaned.strip()
    elif level == 2:
        return re.sub(r'^\s*\d+\.\d+[\.:\s-]*', '', title).strip()
    elif level == 3:
        return re.sub(r'^\s*\d+\.\d+\.\d+[\.:\s-]*', '', title).strip()
    elif level == 4:
        return re.sub(r'^\s*\d+\.\d+\.\d+\.\d+[\.:\s-]*', '', title).strip()
    return title


def set_cell_shading(cell, fill='D9EAF7'):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


# Nhãn caption đúng chữ Word tiếng Việt (NFD), khớp TOC \c "Hình" / "Bảng" trong mẫu.
CAPTION_LABELS = {
    'figure': 'Hi\u0300nh',  # Hình
    'table': 'Ba\u0309ng',    # Bảng
}


def add_field(paragraph, instruction, cached_text='Mở trong Microsoft Word và cập nhật trường để hiển thị.'):
    """Insert a Word field as separate runs, matching native Caption/TOC markup."""
    r = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    r._r.append(begin)

    r = paragraph.add_run()
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction if instruction.startswith(' ') else f' {instruction} '
    r._r.append(instr)

    r = paragraph.add_run()
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    r._r.append(separate)

    r = paragraph.add_run(cached_text)
    rPr = r._r.get_or_add_rPr()
    if rPr.find(qn('w:noProof')) is None:
        rPr.append(OxmlElement('w:noProof'))

    r = paragraph.add_run()
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    r._r.append(end)


def next_bookmark_id(doc):
    max_id = 0
    for el in doc.element.iter(qn('w:bookmarkStart')):
        try:
            max_id = max(max_id, int(el.get(qn('w:id'), '0')))
        except ValueError:
            pass
    return max_id + 1


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); rPr.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); rPr.append(underline)
    r.append(rPr); t = OxmlElement('w:t'); t.text = text; r.append(t); hyperlink.append(r)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, citation_numbers):
    """Small inline Markdown renderer for emphasis, code, links and citations."""
    # Replace citations with numeric citations in first-appearance order.
    def citation_sub(match):
        key = match.group(1)
        if key not in citation_numbers:
            citation_numbers[key] = len(citation_numbers) + 1
        return f'[{citation_numbers[key]}]'
    text = CITATION_RE.sub(citation_sub, text)
    # Tokenize links, bold, italic, and inline code; preserve unmatched text literally.
    pattern = re.compile(r'(\*\*.+?\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos: paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith('**'):
            r = paragraph.add_run(token[2:-2]); r.bold = True
        elif token.startswith('*'):
            r = paragraph.add_run(token[1:-1]); r.italic = True
        elif token.startswith('`'):
            r = paragraph.add_run(token[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(10)
        else:
            lm = LINK_RE.fullmatch(token)
            if lm: add_hyperlink(paragraph, lm.group(1), lm.group(2))
            else: paragraph.add_run(token)
        pos = m.end()
    if pos < len(text): paragraph.add_run(text[pos:])


def add_caption(doc, text, kind, chapter, counter):
    """Insert a native Word caption (Caption style + STYLEREF/SEQ) so it appears in Table of Figures/Tables."""
    label = CAPTION_LABELS[kind]
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = (kind == 'table')

    bmid = next_bookmark_id(doc)
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bmid))
    start.set(qn('w:name'), f'_Ref{kind}_{chapter}_{counter}')
    p._p.append(start)

    p.add_run(f'{label} ')
    add_field(p, 'STYLEREF 1 \\s', str(chapter))
    p.add_run('.')
    add_field(p, f'SEQ {label} \\* ARABIC \\s 1', str(counter))

    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bmid))
    p._p.append(end)

    p.add_run(f' {text}')
    return p


def parse_table(lines):
    rows = []
    for line in lines:
        row = [x.strip() for x in line.strip().strip('|').split('|')]
        if row and all(re.fullmatch(r':?-{3,}:?', x.replace(' ', '')) for x in row):
            continue
        rows.append(row)
    width = max((len(r) for r in rows), default=1)
    return [r + [''] * (width-len(r)) for r in rows]


def clear_body(doc):
    body = doc._element.body
    sectPr = body.sectPr
    for child in list(body):
        if child is not sectPr: body.remove(child)


def build(markdown_files, template, output, refs_file):
    doc = Document(template)
    has_front_matter = len(doc.paragraphs) > 0
    if not has_front_matter:
        clear_body(doc)
    has_numbering = template_has_heading_numbering(doc)
    citation_numbers = {}
    refs = json.loads(Path(refs_file).read_text(encoding='utf-8')) if Path(refs_file).exists() else {}
    chapter = 0
    fig_count = 0
    table_count = 0
    pending_table_caption = None
    used_captions = []

    # If the template does not already contain front matter pages, create a default TOC page.
    if not has_front_matter:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before = False
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(18)
        p.paragraph_format.keep_with_next = True
        r = p.add_run('MỤC LỤC')
        r.bold = True
        r.font.size = Pt(16)
        r.font.name = 'Times New Roman'
        toc = doc.add_paragraph()
        add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
        doc.add_page_break()

    for md_path in markdown_files:
        md_path = Path(md_path)
        lines = md_path.read_text(encoding='utf-8-sig').splitlines()
        i = 0
        while i < len(lines):
            raw = lines[i]; line = raw.strip()
            if not line:
                i += 1; continue
            if line.startswith('```'):
                code_lines=[]; i += 1
                while i < len(lines) and not lines[i].strip().startswith('```'):
                    code_lines.append(lines[i]); i += 1
                if i < len(lines): i += 1
                p = doc.add_paragraph(style='Normal')
                p.paragraph_format.left_indent = Inches(0.25)
                r = p.add_run('\n'.join(code_lines)); r.font.name='Consolas'; r.font.size=Pt(9)
                continue
            hm = HEADING_RE.match(line)
            if hm:
                level = len(hm.group(1)); title = hm.group(2).strip()
                if level == 1:
                    cm = re.search(r'CHƯƠNG\s+(\d+)', title, re.IGNORECASE)
                    if cm: chapter = int(cm.group(1))
                    else:
                        nm = re.match(r'(\d+)\.', title)
                        if nm: chapter = int(nm.group(1))
                    fig_count = table_count = 0
                p = doc.add_paragraph(style=f'Heading {level}')
                display_title = clean_heading_text(title, level) if has_numbering else title
                add_inline(p, display_title, citation_numbers)
                i += 1; continue
            im = IMAGE_RE.match(line)
            if im:
                alt, rel = im.group(1), im.group(2)
                img_path = (md_path.parent / rel).resolve()
                if not img_path.exists():
                    raise FileNotFoundError(f'Ảnh không tồn tại: {img_path} (trong {md_path.name})')
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Mẫu ĐATN 2019: Chiều rộng hình vẽ không quá 75% chiều rộng nội dung phần chữ
                # Chiều rộng nội dung = 21cm - 3.5cm - 2.5cm = 15cm -> 75% = 11.25cm = 4.43 inches
                try: p.add_run().add_picture(str(img_path), width=Inches(4.43))
                except Exception as e: raise RuntimeError(f'Không chèn được ảnh {img_path}: {e}')
                fig_count += 1
                add_caption(doc, alt or img_path.stem.replace('_',' '), 'figure', chapter or 1, fig_count)
                used_captions.append(('figure', chapter or 1, fig_count))
                i += 1; continue
            if line.lower().startswith('table:'):
                pending_table_caption = line.split(':',1)[1].strip(); i += 1; continue
            if '|' in line and i+1 < len(lines) and re.match(r'^\s*\|?\s*:?-{3,}', lines[i+1]):
                table_lines=[]
                while i < len(lines) and '|' in lines[i].strip() and lines[i].strip():
                    table_lines.append(lines[i]); i += 1
                rows = parse_table(table_lines)
                if rows:
                    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                    table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for ri, row in enumerate(rows):
                        for ci, value in enumerate(row):
                            cell = table.cell(ri,ci); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                            cell.text = ''
                            p = cell.paragraphs[0]
                            add_inline(p, value, citation_numbers)
                            if ri == 0:
                                set_cell_shading(cell)
                                for run in p.runs: run.bold = True
                    if pending_table_caption:
                        table_count += 1
                        # Caption goes before the table, consistent with common thesis conventions.
                        table._element.addprevious(add_caption(doc, pending_table_caption, 'table', chapter or 1, table_count)._p)
                        used_captions.append(('table', chapter or 1, table_count))
                        pending_table_caption = None
                continue
            if re.match(r'^\s*[-*+]\s+', raw):
                p = doc.add_paragraph(style='Normal')
                p.paragraph_format.left_indent = Inches(0.25)
                p.add_run('• ')
                add_inline(p, re.sub(r'^\s*[-*+]\s+', '', raw), citation_numbers)
                i += 1; continue
            if re.match(r'^\s*\d+[.)]\s+', raw):
                p = doc.add_paragraph(style='Normal')
                p.paragraph_format.left_indent = Inches(0.25)
                p.add_run(re.match(r'^\s*(\d+)[.)]', raw).group(1) + '. ')
                add_inline(p, re.sub(r'^\s*\d+[.)]\s+', '', raw), citation_numbers)
                i += 1; continue
            # Join consecutive ordinary lines into one paragraph, until another block begins.
            para = [line]; i += 1
            while i < len(lines):
                nxt = lines[i].strip()
                if (not nxt or HEADING_RE.match(nxt) or IMAGE_RE.match(nxt) or nxt.startswith('```') or nxt.lower().startswith('table:') or re.match(r'^\s*[-*+]\s+', lines[i]) or re.match(r'^\s*\d+[.)]\s+', lines[i])):
                    break
                if '|' in nxt and i+1 < len(lines) and re.match(r'^\s*\|?\s*:?-{3,}', lines[i+1]): break
                para.append(nxt); i += 1
            p = doc.add_paragraph(style='Normal')
            add_inline(p, ' '.join(para), citation_numbers)

    # Automatically append bibliography for cited keys in order of first citation.
    if citation_numbers:
        doc.add_page_break()
        p = doc.add_paragraph('TÀI LIỆU THAM KHẢO', style='Heading 1')
        disable_numbering(p)
        reverse = sorted(citation_numbers.items(), key=lambda item: item[1])
        for key, number in reverse:
            item = refs.get(key)
            p = doc.add_paragraph(style='Normal')
            if not item:
                p.add_run(f'[{number}] {key} — CHƯA CÓ THÔNG TIN TRONG references.json').bold = True
                continue
            authors = item.get('authors','')
            year = item.get('year','n.d.')
            title = item.get('title',key)
            publisher = item.get('publisher','')
            url = item.get('url','')
            text = f'[{number}] {authors} ({year}). {title}.'
            if publisher: text += f' {publisher}.'
            add_inline(p, text, citation_numbers)
            if url: p.add_run(' '); add_hyperlink(p, url, url)

    # Set TOC and other Word fields to update on open.
    settings = doc.settings._element
    update = settings.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields'); settings.append(update)
    update.set(qn('w:val'), 'true')
    doc.core_properties.title = Path(output).stem.replace('_',' ')
    doc.save(output)
    print(f'Đã xuất: {Path(output).resolve()}')
    print(f'Số nguồn được trích dẫn: {len(citation_numbers)}')
    print('Lưu ý: mở file bằng Word, nhấn Ctrl+A rồi F9 để cập nhật mục lục/số trang.')


def main():
    parser = argparse.ArgumentParser(description='Xuất Markdown sang Word theo template.')
    parser.add_argument('--template', default='template_clean.docx')
    parser.add_argument('--output', default='output.docx')
    parser.add_argument('--refs', default='references.json')
    parser.add_argument('markdown', nargs='*', help='Các file .md theo thứ tự; mặc định dùng chapters/*.md')
    args = parser.parse_args()
    files = [Path(x) for x in args.markdown] if args.markdown else sorted(Path('chapters').glob('*.md'))
    if not files: sys.exit('Không tìm thấy file Markdown. Hãy thêm chapters/chapter1.md hoặc truyền đường dẫn .md.')
    if not Path(args.template).exists(): sys.exit(f'Không tìm thấy template: {args.template}. Chạy make_template.py trước.')
    build(files, args.template, args.output, args.refs)

if __name__ == '__main__': main()
