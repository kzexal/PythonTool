from docx import Document
from docx.oxml.ns import qn
from lxml import etree

def dump(path):
    print('====', path)
    doc = Document(path)
    print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables), 'sections', len(doc.sections))
    body = doc.element.body
    children = list(body)
    print('body children', len(children), 'tags:', [c.tag.split('}')[-1] for c in children])

    print('\n-- styles of interest --')
    for name in ('Heading 1', 'Heading 2', 'TOC Heading', 'TOCHeading', 'Caption', 'Normal'):
        try:
            s = doc.styles[name]
        except KeyError:
            print(name, 'MISSING')
            continue
        pf = s.paragraph_format
        print(name, 'id', s.style_id, 'based', s.base_style.name if s.base_style else None,
              'page_break_before', pf.page_break_before, 'keep_with_next', pf.keep_with_next,
              'align', pf.alignment)
        pPr = s.element.pPr
        if pPr is not None:
            print('  pPr snippet', etree.tostring(pPr, encoding='unicode')[:400])

    print('\n-- all paragraphs --')
    for i, p in enumerate(doc.paragraphs):
        xml = p._p.xml
        has_page = 'w:type="page"' in xml or "w:type='page'" in xml
        has_br = '<w:br' in xml
        t = p.text.replace('\n', '\\n')
        print(f'{i:3d} style={p.style.name!r:20} pageBR={has_page!s:5} br={has_br!s:5} align={p.alignment} text={t[:90]!r}')

    print('\n-- body child detail --')
    for i, c in enumerate(children):
        tag = c.tag.split('}')[-1]
        texts = ''.join(c.itertext())
        texts = ' '.join(texts.split())
        extra = ''
        if tag == 'p':
            extra = ' pageBR' if 'w:type="page"' in c.xml else ''
        if tag == 'sdt':
            extra = ' SDT ' + ' '.join(c.itertext())[:40]
        print(f'{i:3d} {tag:8}{extra} {texts[:80]!r}')

    # docGrid
    for i, sect in enumerate(doc.sections):
        sp = sect._sectPr
        print('\nsection', i, etree.tostring(sp, encoding='unicode')[:800])

for path in ('Mẫu ĐATN_2019_version 1_1.docx', 'template_clean.docx', 'Bao_cao_do_an.docx'):
    try:
        dump(path)
    except Exception as e:
        print('FAIL', path, e)
    print('\n\n')
