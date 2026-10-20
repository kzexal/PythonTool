"""
md2pptx.py - Chuyển file Markdown thành PowerPoint theo template HUST.

QUY ƯỚC MARKDOWN
----------------
## Tên chương        -> slide ngăn cách chương (nền xanh đậm, chữ trắng). Cũng là mục trong "Nội dung chính".
# Tiêu đề slide       -> một slide nội dung (tiêu đề nằm ở thanh trên cùng).
- Nội dung            -> bullet cấp 1 của slide hiện tại.
  - Nội dung phụ      -> bullet cấp 2 (thụt 2 dấu cách hoặc 1 tab).
[chữ](https://...)    -> hyperlink bấm được (dùng được trong bullet).

Cách dùng:
    python md2pptx.py                      # dùng input.md, output.pptx, template mặc định
    python md2pptx.py in.md out.pptx tpl.pptx
"""
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "Arial"
SECTION_COLOR = RGBColor(0, 51, 102)
WHITE = RGBColor(255, 255, 255)
PAGE_NUM_COLOR = RGBColor(190, 30, 60)   # đỏ cùng tông với đường kẻ của template
PAGE_NUM_TOP = 7.1                       # inch; đặt số trang dưới đường kẻ đỏ (chỉnh nếu cần)

link_pattern = re.compile(r'\[([^\]]+)\]\((https?://[^\)]+)\)')


# ----------------------------------------------------------------------
# Làm sạch input
# ----------------------------------------------------------------------
def normalize(line):
    """Dọn các lỗi thường gặp khi file .md bị export/escape sai:
       - tiền tố '# ' thừa đứng trước dòng đã bị escape (vd '# \\## Phần 1')
       - '&#x20;' thay cho dấu cách
       - ký tự escape bằng backslash (\\#, \\-, \\[ ...)
    """
    s = line.rstrip('\r\n')

    if s.strip() == '#':                      # dòng "#" trống
        return ''
    if re.match(r'^# (?=[\\&])', s):          # "# \..." hoặc "# &#x20;..."
        s = s[2:]

    s = s.replace('&#x20;', ' ').replace('&nbsp;', ' ').replace('\u00a0', ' ')
    s = re.sub(r'\\([\\`*_{}\[\]()#+\-.!])', r'\1', s)
    return s


def parse_markdown(md_filepath):
    """Trả về (slides_data, chapters)."""
    with open(md_filepath, 'r', encoding='utf-8-sig') as f:
        lines = f.readlines()

    slides_data, chapters = [], []

    for line in lines:
        s = normalize(line)
        if not s.strip():
            continue

        # Thụt lề: tab -> 2 dấu cách để dễ nhận bullet cấp 2
        s = s.replace('\t', '  ')
        stripped = s.strip()
        indent = len(s) - len(s.lstrip(' '))

        # QUY TẮC 1: "## " -> tên chương
        if s.startswith('## '):
            title = s[3:].strip()
            chapters.append(title)
            slides_data.append({'type': 'section', 'title': title})

        # QUY TẮC 2: "# " -> tiêu đề slide
        elif s.startswith('# '):
            slides_data.append({'type': 'content',
                                'title': s[2:].strip(),
                                'content': []})

        # QUY TẮC 3 & 4: bullet "- " hoặc "* "
        elif stripped.startswith(('- ', '* ')):
            if slides_data and slides_data[-1]['type'] == 'content':
                level = 1 if indent >= 2 else 0
                slides_data[-1]['content'].append(
                    {'level': level, 'text': stripped[2:].strip()})
        # Các dòng khác bị bỏ qua

    return slides_data, chapters


# ----------------------------------------------------------------------
# Hàm hỗ trợ định dạng
# ----------------------------------------------------------------------
def set_font(paragraph, font_name=FONT, font_size_pt=24, bold=False,
             center=False, color=None):
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = Pt(font_size_pt)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color
    if center:
        paragraph.alignment = PP_ALIGN.CENTER


def set_indent(paragraph, left_inches, bullet_char=None):
    """Thụt lề thật (marL/indent) và bullet thật của PowerPoint."""
    pPr = paragraph._p.get_or_add_pPr()
    hang = Inches(0.4)
    pPr.set('marL', str(int(Inches(left_inches) + hang)))
    pPr.set('indent', str(-int(hang)))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if bullet_char:
        bu = pPr.makeelement(qn('a:buChar'), {'char': bullet_char})
        pPr.append(bu)


def clear_all_placeholders(slide):
    for shape in list(slide.shapes):
        if shape.is_placeholder:
            el = shape._element
            el.getparent().remove(el)


def create_textbox_title(slide, text):
    txBox = slide.shapes.add_textbox(Inches(0.4), Inches(0), Inches(12.0), Inches(0.7))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = tf.paragraphs[0]
    p.text = text
    set_font(p, FONT, 39, bold=True, color=WHITE)


def append_text_with_hyperlinks(paragraph, text, font_size):
    def add(t, url=None):
        run = paragraph.add_run()
        run.text = t
        run.font.name = FONT
        run.font.size = Pt(font_size)
        if url:
            run.hyperlink.address = url

    last = 0
    for m in link_pattern.finditer(text):
        if m.start() > last:
            add(text[last:m.start()])
        add(m.group(1), m.group(2))
        last = m.end()
    if last < len(text):
        add(text[last:])


def output_is_locked(path):
    """True nếu file đang được mở ở chương trình khác (vd PowerPoint) nên không ghi được."""
    if not os.path.exists(path):
        return False
    try:
        with open(path, 'ab'):
            return False
    except PermissionError:
        return True


def remove_slide(prs, slide_id_el):
    prs.part.drop_rel(slide_id_el.get(qn('r:id')))
    prs.slides._sldIdLst.remove(slide_id_el)


# ----------------------------------------------------------------------
# Hàm chính
# ----------------------------------------------------------------------
def create_presentation_from_md(md_filepath, output_filepath, template_filepath):
    if output_is_locked(output_filepath):
        print(f"Lỗi: File '{output_filepath}' đang được mở trong PowerPoint (hoặc chương trình khác).\n"
              f"Hãy TẮT file đó đi rồi chạy lại lệnh.")
        return False

    prs = Presentation(template_filepath)
    slides_data, chapters = parse_markdown(md_filepath)

    n_template = len(prs.slides)          # số slide có sẵn trong template
    layout = prs.slide_layouts[1]

    # ---- Trang bìa (slide đầu của template) ----
    title_slide = prs.slides[0]
    if title_slide.shapes.title:
        title_slide.shapes.title.text = "TIÊU ĐỀ BÀI THUYẾT TRÌNH CỦA BẠN"
        title_slide.shapes.title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        set_font(title_slide.shapes.title.text_frame.paragraphs[0],
                 FONT, 44, bold=True, center=True)

    if len(title_slide.shapes.placeholders) > 1:
        sub = title_slide.shapes.placeholders[1]
        sub.text = "Nhóm 1 - Lớp ABC\nGiảng viên hướng dẫn: TS. Nguyễn Văn A"
        for p in sub.text_frame.paragraphs:
            set_font(p, FONT, 24, bold=False, center=True)

    # ---- Slide Mục lục: liệt kê các CHƯƠNG (##) ----
    if chapters:
        toc = prs.slides.add_slide(layout)
        clear_all_placeholders(toc)
        create_textbox_title(toc, "Nội dung chính")

        tf = toc.shapes.add_textbox(Inches(1.2), Inches(1.3),
                                    Inches(10), Inches(5.5)).text_frame
        tf.word_wrap = True
        for i, chap in enumerate(chapters):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = chap
            set_font(p, FONT, 32 if len(chapters) <= 6 else 24, bold=True)
            p.space_after = Pt(36 if len(chapters) <= 6 else 16)

    # ---- Các slide chương + nội dung ----
    for info in slides_data:
        if info['type'] == 'section':
            s = prs.slides.add_slide(layout)
            clear_all_placeholders(s)

            bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                    prs.slide_width, prs.slide_height)
            bg.fill.solid()
            bg.fill.fore_color.rgb = SECTION_COLOR
            bg.line.color.rgb = SECTION_COLOR

            tf = s.shapes.add_textbox(Inches(0.5), Inches(2.5),
                                      Inches(12.33), Inches(2.5)).text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = info['title']
            set_font(p, FONT, 54, bold=True, center=True, color=WHITE)

        else:  # content
            s = prs.slides.add_slide(layout)
            clear_all_placeholders(s)
            create_textbox_title(s, info['title'])

            items = info['content']
            if not items:
                continue

            tf = s.shapes.add_textbox(Inches(0.8), Inches(1.3),
                                      Inches(11.7), Inches(5.5)).text_frame
            tf.word_wrap = True

            n = len(items)
            if n <= 3:
                size0, size1, sb, sa = 32, 28, 12, 24
            elif n > 6:
                size0, size1, sb, sa = 24, 20, 6, 8
            else:
                size0, size1, sb, sa = 28, 24, 8, 12

            for i, item in enumerate(items):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                if item['level'] == 0:
                    size = size0
                    set_indent(p, 0.0, "•")
                    p.space_before = Pt(sb)
                    p.space_after = Pt(sa)
                else:
                    size = size1
                    set_indent(p, 0.6, "◦")
                    p.space_after = Pt(sa / 2)
                append_text_with_hyperlinks(p, item['text'], size)

    # ---- Xếp lại slide ----
    # Template: slide 0 = bìa, slide cuối = kết thúc, các slide giữa = mẫu (xoá đi).
    sld_ids = list(prs.slides._sldIdLst)
    has_closing = n_template >= 2
    if has_closing:
        closing = sld_ids[n_template - 1]
        for el in sld_ids[1:n_template - 1]:
            remove_slide(prs, el)
        prs.slides._sldIdLst.remove(closing)
        prs.slides._sldIdLst.append(closing)

    # ---- Đánh số trang (bỏ bìa và slide kết thúc) ----
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides):
        if i == 0 or (has_closing and i == total - 1):
            continue
        tf = slide.shapes.add_textbox(Inches(12.0), Inches(PAGE_NUM_TOP),
                                      Inches(1), Inches(0.3)).text_frame
        tf.margin_bottom = tf.margin_top = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.text = str(i)
        set_font(p, FONT, 16, bold=True, color=PAGE_NUM_COLOR)

    try:
        prs.save(output_filepath)
    except PermissionError:
        print(f"Lỗi: Không ghi được '{output_filepath}' vì file đang được mở.\n"
              f"Hãy TẮT file đó đi rồi chạy lại lệnh.")
        return False
    print(f"Đã xuất bản file PPTX hoàn chỉnh tại: {output_filepath}")
    return True


if __name__ == "__main__":
    args = sys.argv[1:]
    input_md = args[0] if len(args) > 0 else "input.md"
    output_pptx = args[1] if len(args) > 1 else "output.pptx"
    template_pptx = args[2] if len(args) > 2 else "HUST_PPT_template_2022_Blue_16x9.pptx"

    if os.path.exists(input_md) and os.path.exists(template_pptx):
        create_presentation_from_md(input_md, output_pptx, template_pptx)
    else:
        print("Lỗi: Không tìm thấy file input.md hoặc file template!")
