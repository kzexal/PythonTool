
from pathlib import Path
from datetime import datetime
import hashlib
import os
import re
import tempfile
import unicodedata

import genanki
from openpyxl import load_workbook


# =====================================================
# 1. CAU HINH DU AN
# =====================================================

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "input" / "vocabulary.xlsx"
OUTPUT_DIR = BASE_DIR / "output"

OUTPUT_FILE = OUTPUT_DIR / "Japanese_Vocabulary.apkg"
REPORT_FILE = OUTPUT_DIR / "import_report.txt"

DECK_NAME = "Japanese Vocabulary"

# Giu cac ID nay co dinh giua cac lan chay.
# Neu da co bo Anki cu, khong doi ID tuy tien.
DECK_ID = 2059400110
MODEL_ID = 1607392319

REQUIRED_COLUMNS = ["Japanese", "Vietnamese"]

OPTIONAL_COLUMNS = [
    "Reading",
    "Example",
    "ExampleMeaning",
    "JLPT",
    "Notes",
    "Tags",
]

VALID_JLPT = {"N5", "N4", "N3", "N2", "N1"}


# =====================================================
# 2. HAM CHUAN HOA DU LIEU
# =====================================================

def clean(value):
    if value is None:
        return ""

    text = unicodedata.normalize("NFC", str(value))
    return text.strip()


def normalize(value):
    """Chuan hoa Unicode va khoang trang de so sanh."""
    text = clean(value)
    return re.sub(r"\s+", " ", text)


def html_escape(value):
    """Dam bao noi dung Excel khong lam hong HTML cua the."""
    text = clean(value)

    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#x27;")
    )


def parse_tags(value):
    """Chap nhan tag phan cach bang dau phay hoac cham phay."""
    parts = re.split(r"[,;]", clean(value))
    tags = []

    for part in parts:
        tag = part.strip()

        if tag:
            # Anki tags khong nen chua khoang trang.
            tag = re.sub(r"\s+", "_", tag)
            if tag not in tags:
                tags.append(tag)

    return tags


# =====================================================
# 3. DOC VA KIEM TRA EXCEL
# =====================================================

def read_excel():
    if not INPUT_FILE.is_file():
        raise FileNotFoundError(
            f"Khong tim thay file Excel:\n{INPUT_FILE}"
        )

    workbook = load_workbook(
        INPUT_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)

        if not header_row:
            raise ValueError("Sheet Excel dang trong.")

        headers = [clean(value) for value in header_row]

        header_index = {}

        for index, header in enumerate(headers):
            if not header:
                continue

            key = header.casefold()

            if key in header_index:
                raise ValueError(
                    f"Ten cot bi trung: {header}"
                )

            header_index[key] = index

        missing = [
            name for name in REQUIRED_COLUMNS
            if name.casefold() not in header_index
        ]

        if missing:
            raise ValueError(
                "Thieu cot bat buoc: " + ", ".join(missing)
            )

        words = []
        issues = []
        seen = {}
        duplicate_count = 0
        invalid_count = 0

        for row_number, row in enumerate(iterator, start=2):

            def get_value(name):
                index = header_index.get(name.casefold())

                if index is None or index >= len(row):
                    return ""

                return clean(row[index])

            # Bo qua dong trong hoan toan
            if not any(clean(value) for value in row):
                continue

            word = {
                "Japanese": get_value("Japanese"),
                "Reading": get_value("Reading"),
                "Vietnamese": get_value("Vietnamese"),
                "Example": get_value("Example"),
                "ExampleMeaning": get_value("ExampleMeaning"),
                "JLPT": get_value("JLPT").upper(),
                "Notes": get_value("Notes"),
                "Tags": parse_tags(get_value("Tags")),
                "row_number": row_number,
            }

            japanese = word["Japanese"]
            vietnamese = word["Vietnamese"]

            # Loi: thieu du lieu bat buoc
            if not japanese or not vietnamese:
                issues.append(
                    f"[LOI] Dong {row_number}: "
                    "thieu Japanese hoac Vietnamese."
                )
                invalid_count += 1
                continue

            # Kiem tra cap do JLPT
            jlpt = word["JLPT"]

            if jlpt and jlpt not in VALID_JLPT:
                issues.append(
                    f"[CANH BAO] Dong {row_number}: "
                    f"JLPT '{jlpt}' khong hop le."
                )

            # Neu co cau vi du ma chua co ban dich
            if word["Example"] and not word["ExampleMeaning"]:
                issues.append(
                    f"[CANH BAO] Dong {row_number}: "
                    "co Example nhung thieu ExampleMeaning."
                )

            # Dinh danh tu: Japanese + Reading.
            # Neu cung tu va cach doc nhung khac nghia,
            # can xem lai vi ban nay chua quan ly sense ID.
            identity = (
                normalize(japanese),
                normalize(word["Reading"]),
            )

            if identity in seen:
                old = seen[identity]

                if normalize(old["Vietnamese"]) == normalize(vietnamese):
                    duplicate_count += 1
                    issues.append(
                        f"[TRUNG] Dong {row_number}: "
                        f"'{japanese}' da co o dong "
                        f"{old['row_number']}; bo qua."
                    )
                else:
                    issues.append(
                        f"[CAN KIEM TRA] Dong {row_number}: "
                        f"'{japanese}' co cung cach doc nhung "
                        "khac nghia voi dong "
                        f"{old['row_number']}; bo qua de tranh "
                        "ghi de nham ghi chu Anki."
                    )

                invalid_count += 1
                continue

            seen[identity] = word
            words.append(word)

        return words, issues, duplicate_count, invalid_count

    finally:
        workbook.close()


# =====================================================
# 4. DINH NGHIA MAU NOTE VA FLASHCARD
# =====================================================

def create_model():
    return genanki.Model(
        MODEL_ID,
        "Japanese Vocabulary Model",
        fields=[
            {"name": "Japanese"},
            {"name": "Reading"},
            {"name": "Vietnamese"},
            {"name": "Example"},
            {"name": "ExampleMeaning"},
            {"name": "Notes"},
        ],
        templates=[
            {
                "name": "Japanese to Vietnamese",
                "qfmt": """
                    <div class="jp">{{Japanese}}</div>
                """,
                "afmt": """
                    {{FrontSide}}
                    <hr id="answer">
                    <div class="reading">{{Reading}}</div>
                    <div class="meaning">{{Vietnamese}}</div>
                    {{#Example}}
                    <div class="example">{{Example}}</div>
                    {{/Example}}
                    {{#ExampleMeaning}}
                    <div class="example-meaning">
                        {{ExampleMeaning}}
                    </div>
                    {{/ExampleMeaning}}
                    {{#Notes}}
                    <div class="notes">{{Notes}}</div>
                    {{/Notes}}
                """,
            },
            {
                "name": "Vietnamese to Japanese",
                "qfmt": """
                    <div class="meaning">{{Vietnamese}}</div>
                """,
                "afmt": """
                    {{FrontSide}}
                    <hr id="answer">
                    <div class="jp">{{Japanese}}</div>
                    <div class="reading">{{Reading}}</div>
                    {{#Example}}
                    <div class="example">{{Example}}</div>
                    {{/Example}}
                    {{#ExampleMeaning}}
                    <div class="example-meaning">
                        {{ExampleMeaning}}
                    </div>
                    {{/ExampleMeaning}}
                    {{#Notes}}
                    <div class="notes">{{Notes}}</div>
                    {{/Notes}}
                """,
            },
        ],
        css="""
            .card {
                font-family: Arial, sans-serif;
                font-size: 20px;
                text-align: center;
                color: #222;
                background: #fff;
                padding: 20px;
            }

            .jp {
                font-size: 42px;
                margin: 20px 0;
            }

            .reading {
                font-size: 24px;
                color: #168a85;
                margin: 12px 0;
            }

            .meaning {
                font-size: 26px;
                margin: 16px 0;
            }

            .example {
                margin-top: 24px;
                font-size: 20px;
            }

            .example-meaning {
                margin-top: 8px;
                color: #666;
                font-size: 16px;
            }

            .notes {
                margin-top: 18px;
                color: #777;
                font-size: 15px;
            }
        """,
    )


# =====================================================
# 5. TAO DECK
# =====================================================

def create_deck(words):
    model = create_model()
    deck = genanki.Deck(DECK_ID, DECK_NAME)

    for word in words:
        fields = [
            html_escape(word["Japanese"]),
            html_escape(word["Reading"]),
            html_escape(word["Vietnamese"]),
            html_escape(word["Example"]),
            html_escape(word["ExampleMeaning"]),
            html_escape(word["Notes"]),
        ]

        # ID ghi chu khong thay doi khi sua nghia/vi du.
        # Can giu Japanese va Reading on dinh.
        note_guid = genanki.guid_for(
            normalize(word["Japanese"]),
            normalize(word["Reading"]),
        )

        note = genanki.Note(
            model=model,
            fields=fields,
            guid=note_guid,
            tags=word["Tags"],
        )

        deck.add_note(note)

    return deck


# =====================================================
# 6. GHI BAO CAO
# =====================================================

def write_report(
    words,
    issues,
    duplicate_count,
    invalid_count,
    export_status,
):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    lines = [
        "JAPANESE FLASHCARDS - IMPORT REPORT",
        f"Thoi gian: {datetime.now():%Y-%m-%d %H:%M:%S}",
        f"File nguon: {INPUT_FILE}",
        f"Trang thai xuat: {export_status}",
        f"So tu hop le: {len(words)}",
        f"So the du kien: {len(words) * 2}",
        f"So dong trung: {duplicate_count}",
        f"So dong bi bo qua: {invalid_count}",
        f"So canh bao/loi: {len(issues)}",
        "",
        "CHI TIET:",
    ]

    lines.extend(issues or ["Khong co loi hoac canh bao."])

    REPORT_FILE.write_text(
        "\n".join(lines),
        encoding="utf-8",
    )


# =====================================================
# 7. XUAT APKG AN TOAN
# =====================================================

def export_apkg(deck):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # Tao file tam ngay trong output de os.replace
    # co the thay the file dich tren cung volume.
    file_descriptor, temp_name = tempfile.mkstemp(
        prefix="anki_temp_",
        suffix=".apkg",
        dir=OUTPUT_DIR,
    )
    os.close(file_descriptor)

    temp_path = Path(temp_name)

    try:
        genanki.Package(deck).write_to_file(str(temp_path))

        if not temp_path.is_file() or temp_path.stat().st_size == 0:
            raise OSError("File APKG tao ra rong hoac khong ton tai.")

        # Chi thay file cu sau khi tao file moi thanh cong
        os.replace(temp_path, OUTPUT_FILE)

    finally:
        if temp_path.exists():
            temp_path.unlink()


# =====================================================
# 8. CHUONG TRINH CHINH
# =====================================================

def main():
    words = []
    issues = []
    duplicate_count = 0
    invalid_count = 0

    try:
        words, issues, duplicate_count, invalid_count = read_excel()

        if not words:
            write_report(
                words,
                issues,
                duplicate_count,
                invalid_count,
                "KHONG XUAT - KHONG CO TU HOP LE",
            )
            print("Khong co tu hop le. Hay xem import_report.txt.")
            return

        deck = create_deck(words)
        export_apkg(deck)

        write_report(
            words,
            issues,
            duplicate_count,
            invalid_count,
            "THANH CONG",
        )

        print("\n=== KET QUA ===")
        print(f"Tu hop le: {len(words)}")
        print(f"The du kien: {len(words) * 2}")
        print(f"Dong trung: {duplicate_count}")
        print(f"Dong bi bo qua: {invalid_count}")
        print(f"APKG: {OUTPUT_FILE}")
        print(f"Bao cao: {REPORT_FILE}")

        if issues:
            print("Co canh bao. Hay xem import_report.txt.")

    except Exception as error:
        message = f"[LOI XU LY] {type(error).__name__}: {error}"
        print(message)

        # Neu xuat loi, khong thay the APKG cu.
        # Co gang ghi lai loi vao bao cao.
        try:
            OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
            REPORT_FILE.write_text(
                f"Thoi gian: {datetime.now():%Y-%m-%d %H:%M:%S}\n"
                f"{message}\n",
                encoding="utf-8",
            )
        except OSError:
            pass


if __name__ == "__main__":
    main()
