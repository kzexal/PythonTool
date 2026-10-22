
from pathlib import Path
from datetime import datetime
import shutil

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "input" / "vocabulary.xlsx"


SHEETS = {
    "Vocabulary": [
        "Japanese",
        "Reading",
        "Vietnamese",
        "Example",
        "ExampleMeaning",
        "JLPT",
        "Notes",
        "Tags",
    ],
    "Grammar": [
        "Grammar",
        "Meaning",
        "Structure",
        "Example",
        "ExampleMeaning",
        "JLPT",
        "Notes",
        "Tags",
    ],
    "Kanji": [
        "Kanji",
        "Onyomi",
        "Kunyomi",
        "Meaning",
        "Radical",
        "StrokeCount",
        "ExampleWords",
        "Notes",
        "JLPT",
        "Tags",
    ],
    "Settings": [
        "Key",
        "Value",
        "Description",
    ],
}


DEFAULT_SETTINGS = [
    (
        "AnkiConnectUrl",
        "http://127.0.0.1:8765",
        "Dia chi AnkiConnect",
    ),
    (
        "VocabularyDeck",
        "Japanese Vocabulary",
        "Deck tu vung",
    ),
    (
        "VocabularyModel",
        "Japanese Vocabulary Model",
        "Note type tu vung hien tai",
    ),
    (
        "GrammarDeck",
        "Japanese Grammar",
        "Deck ngu phap",
    ),
    (
        "GrammarModel",
        "Japanese Grammar Model",
        "Note type ngu phap",
    ),
    (
        "KanjiDeck",
        "Japanese Kanji",
        "Deck Kanji",
    ),
    (
        "KanjiModel",
        "Japanese Kanji Model",
        "Note type Kanji",
    ),
    (
        "Version",
        "1.0",
        "Phien ban workbook",
    ),
]


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 24

    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="285A72",
        )
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    for column in ws.columns:
        values = [
            str(cell.value or "")
            for cell in list(column)[:30]
        ]

        width = min(
            max(
                max(
                    (len(value) for value in values),
                    default=10,
                ) + 2,
                12,
            ),
            32,
        )

        ws.column_dimensions[
            get_column_letter(column[0].column)
        ].width = width


def add_missing_headers(ws, headers):
    current_headers = [
        str(cell.value or "").strip().casefold()
        for cell in ws[1]
    ]

    for header in headers:
        if header.casefold() not in current_headers:
            ws.cell(
                row=1,
                column=ws.max_column + 1,
                value=header,
            )
            current_headers.append(header.casefold())


def main():
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Khong tim thay file: {EXCEL_FILE}"
        )

    # Tao ban sao truoc khi sua Excel
    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    backup_file = EXCEL_FILE.with_name(
        f"vocabulary_backup_{timestamp}.xlsx"
    )

    shutil.copy2(EXCEL_FILE, backup_file)

    print(f"Da sao luu: {backup_file}")

    workbook = load_workbook(EXCEL_FILE)

    # Doi ten sheet dau tien thanh Vocabulary
    # neu workbook chua co sheet nay.
    if "Vocabulary" not in workbook.sheetnames:
        workbook.worksheets[0].title = "Vocabulary"

    vocabulary = workbook["Vocabulary"]

    # Chi them cot con thieu, khong xoa du lieu cu.
    add_missing_headers(
        vocabulary,
        SHEETS["Vocabulary"],
    )

    # Tao cac sheet moi.
    for sheet_name in (
        "Grammar",
        "Kanji",
        "Settings",
    ):
        if sheet_name not in workbook.sheetnames:
            ws = workbook.create_sheet(sheet_name)

            for column, header in enumerate(
                SHEETS[sheet_name],
                start=1,
            ):
                ws.cell(
                    row=1,
                    column=column,
                    value=header,
                )

    # Them cot con thieu cho Grammar va Kanji.
    for sheet_name in ("Grammar", "Kanji"):
        add_missing_headers(
            workbook[sheet_name],
            SHEETS[sheet_name],
        )

    # Dien cau hinh mac dinh, khong ghi de key da co.
    settings = workbook["Settings"]

    settings_headers = {
        str(cell.value or "").strip().casefold(): index
        for index, cell in enumerate(
            settings[1],
            start=1,
        )
        if cell.value
    }

    if "key" not in settings_headers or "value" not in settings_headers:
        if settings.max_row > 1:
            raise ValueError(
                "Sheet Settings co du lieu nhung khong dung cau truc. "
                "Hay kiem tra sheet nay thu cong."
            )

        settings.delete_rows(
            1,
            settings.max_row,
        )

        for column, header in enumerate(
            SHEETS["Settings"],
            start=1,
        ):
            settings.cell(
                row=1,
                column=column,
                value=header,
            )

        settings_headers = {
            "key": 1,
            "value": 2,
            "description": 3,
        }

    key_column = settings_headers["key"]
    value_column = settings_headers["value"]
    description_column = settings_headers.get("description", 3)

    existing_keys = {
        str(settings.cell(row, key_column).value or "").strip()
        for row in range(2, settings.max_row + 1)
    }

    for key, value, description in DEFAULT_SETTINGS:
        if key not in existing_keys:
            row = settings.max_row + 1

            settings.cell(row, key_column, key)
            settings.cell(row, value_column, value)
            settings.cell(
                row,
                description_column,
                description,
            )

    # Dinh dang cac sheet
    for ws in workbook.worksheets:
        style_sheet(ws)

    workbook.save(EXCEL_FILE)
    workbook.close()

    print()
    print("NANG CAP THANH CONG!")
    print("Cac sheet:")
    print("- Vocabulary")
    print("- Grammar")
    print("- Kanji")
    print("- Settings")
    print()
    print("Du lieu cu duoc giu lai.")
    print("Hay mo Excel va kiem tra truoc khi tiep tuc.")


if __name__ == "__main__":
    main()
