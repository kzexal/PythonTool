import openpyxl
from openpyxl.styles import Alignment
from pathlib import Path
import pykakasi

# Khởi tạo công cụ chuyển đổi Kanji -> Hiragana
kks = pykakasi.kakasi()

# Đường dẫn tới file Excel
INPUT_FILE = Path("vocabulary.xlsx")
OUTPUT_FILE = Path("vocabulary_updated.xlsx")

def add_reading(text):
    """Hàm tự động thêm cách đọc Hiragana vào sau chữ Kanji"""
    result = kks.convert(text)
    out_text = ""
    for item in result:
        orig = item['orig']
        hira = item['hira']
        
        has_kanji = any('\u4e00' <= char <= '\u9faf' for char in orig)
        
        if has_kanji and orig != hira:
            out_text += f"{orig}({hira})"
        else:
            out_text += orig
            
    return out_text

def update_grammar_sheet():
    if not INPUT_FILE.is_file():
        raise FileNotFoundError(f"Không tìm thấy file Excel:\n{INPUT_FILE}")

    workbook = openpyxl.load_workbook(INPUT_FILE)
    
    if "Grammar" not in workbook.sheetnames:
        raise ValueError("Không tìm thấy sheet Grammar trong file Excel.")

    sheet = workbook["Grammar"]
    headers = [cell.value for cell in sheet[1]]
    
    try:
        col_grammar = headers.index("Grammar")
        col_meaning_usage = headers.index("Meaning")
        col_structure = headers.index("Structure")
        col_example = headers.index("Example")
        col_meaning = headers.index("ExampleMeaning")
        col_notes = headers.index("Notes")
    except ValueError:
        print("Lỗi: Không tìm thấy đủ các cột cần thiết.")
        return

    # Thêm 2 cột mới cho Flashcard
    col_front_idx = len(headers) + 1
    col_back_idx = len(headers) + 2
    sheet.cell(row=1, column=col_front_idx, value="Front")
    sheet.cell(row=1, column=col_back_idx, value="Back")

    # Thuộc tính ép Excel phải hiển thị xuống dòng
    wrap_alignment = Alignment(wrap_text=True)

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=False), start=2):
        grammar_full = row[col_grammar].value
        meaning_usage = row[col_meaning_usage].value 
        structure = row[col_structure].value
        example = row[col_example].value
        example_meaning = row[col_meaning].value
        notes = row[col_notes].value

        if not grammar_full or not example:
            continue

        keyword = grammar_full.split()[0] if " " in grammar_full else grammar_full

        # Tạo mặt trước (Front)
        front_example = example.replace(keyword, " [ ... ] ", 1)
        front_example_with_kana = add_reading(front_example)
        front_content = f"{front_example_with_kana}\n({example_meaning})"
        
        # Tạo mặt sau (Back) với \n để xuống dòng
        back_content = f"**{grammar_full}**\n*Ý nghĩa/Cách dùng:* {meaning_usage}\n*Ngữ pháp:* {structure}"
        if notes:
            back_content += f"\n*Lưu ý:* {notes}"

        # Ghi kết quả
        front_cell = sheet.cell(row=row_idx, column=col_front_idx, value=front_content)
        back_cell = sheet.cell(row=row_idx, column=col_back_idx, value=back_content)
        
        # Cài đặt tự động Wrap Text (Xuống dòng) cho ô
        front_cell.alignment = wrap_alignment
        back_cell.alignment = wrap_alignment

    workbook.save(OUTPUT_FILE)
    print(f"Đã cập nhật xong! File mới được lưu tại: {OUTPUT_FILE}")

if __name__ == "__main__":
    update_grammar_sheet()